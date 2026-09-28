
from io import BytesIO
import os
import re

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF


# ============================================================
# LegalEase - Document Export Service
# DOCX and PDF generation
# ============================================================


def create_docx(document_text: str) -> BytesIO:
    """Convert generated legal text into an editable Word document."""

    doc = Document()

    # Page layout
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    def add_formatted_text(paragraph, text):
        """Support basic Markdown bold formatting."""
        parts = re.split(r"(\*\*.*?\*\*)", text)

        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            else:
                paragraph.add_run(part)

    for raw_line in document_text.splitlines():
        line = raw_line.strip()

        if not line or line == "---" or line == "<br>":
            continue

        if line.startswith("### "):
            paragraph = doc.add_heading(line[4:], level=2)

        elif line.startswith("## "):
            paragraph = doc.add_heading(line[3:], level=1)

        elif line.startswith("# "):
            paragraph = doc.add_heading(line[2:], level=0)

        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_formatted_text(paragraph, line[2:])

        elif re.match(r"^\d+[.)]\s+", line):
            content = re.sub(r"^\d+[.)]\s+", "", line)
            paragraph = doc.add_paragraph(style="List Number")
            add_formatted_text(paragraph, content)

        elif line.startswith("|"):
            paragraph = doc.add_paragraph()
            add_formatted_text(paragraph, line.replace("|", "   "))

        else:
            paragraph = doc.add_paragraph()
            add_formatted_text(paragraph, line)

    # Create Word document in memory
    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output


# ============================================================
# PDF Export
# ============================================================


def create_pdf(document_text: str) -> BytesIO:
    """Convert generated legal text into a PDF document."""

    pdf = FPDF()
    pdf.set_margins(left=20, top=18, right=20)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    # Windows font paths
    font_regular = r"C:\Windows\Fonts\arial.ttf"
    font_bold = r"C:\Windows\Fonts\arialbd.ttf"

    # Use Unicode fonts when available
    if os.path.exists(font_regular) and os.path.exists(font_bold):
        pdf.add_font("ArialUnicode", style="", fname=font_regular)
        pdf.add_font("ArialUnicode", style="B", fname=font_bold)
        font_name = "ArialUnicode"
    else:
        font_name = "Helvetica"

    pdf.set_font(font_name, size=10)

    for raw_line in document_text.splitlines():
        line = raw_line.strip()

        if not line or line == "---" or line == "<br>":
            continue

        # Remove basic Markdown markers
        line = line.replace("**", "").replace("`", "")

        if line.startswith("# "):
            pdf.ln(3)
            pdf.set_font(font_name, style="B", size=17)
            pdf.multi_cell(w=0, h=9, text=line[2:])
            pdf.ln(2)
            pdf.set_font(font_name, size=10)

        elif line.startswith("## "):
            pdf.ln(2)
            pdf.set_font(font_name, style="B", size=13)
            pdf.multi_cell(w=0, h=7, text=line[3:])
            pdf.ln(1)
            pdf.set_font(font_name, size=10)

        elif line.startswith("### "):
            pdf.ln(2)
            pdf.set_font(font_name, style="B", size=11)
            pdf.multi_cell(w=0, h=6, text=line[4:])
            pdf.set_font(font_name, size=10)

        elif line.startswith("- "):
            pdf.multi_cell(w=0, h=6, text="- " + line[2:])

        elif re.match(r"^\d+[.)]\s+", line):
            pdf.multi_cell(w=0, h=6, text=line)

        elif line.startswith("|"):
            # Preserve table content in a readable text format
            pdf.multi_cell(w=0, h=6, text=line.replace("|", "   "))

        else:
            pdf.multi_cell(w=0, h=6, text=line)

    # Save PDF to memory
    pdf_bytes = bytes(pdf.output())
    output = BytesIO(pdf_bytes)
    output.seek(0)

    return output